import numpy as np
import io
import zipfile
from datetime import datetime, timedelta

from openpyxl import load_workbook

from src.gate import GateLog, GateTracker
from src.store import Store, needs_review, to_excel, to_zip

FRAME = np.zeros((100, 200, 3), np.uint8)


def code(c, valid=True, score=0.9):
    return {'code': c, 'valid': valid, 'size_type': '45G1', 'score': score, 'box': (10, 10, 50, 20)}


def plate(digits='3231', letters='UTA', verified=True, score=0.9):
    return {'plate': f'{digits} {letters}', 'digits': digits, 'letters': letters, 'arabic': 'و ط ا ٣٢٣١',
            'verified': verified, 'score': score, 'box': (20, 60, 40, 20)}


def run(tracker, frames):
    """frames: [(now, codes, plates)] -> [(now, event)]"""
    return [(now, e) for now, cs, ps in frames for e in tracker.update(cs, FRAME, now, ps)]


def test_plate_and_container_make_one_event():
    t = GateTracker(min_hits=2)
    events = run(t, [(0, [code('MRSU4425761')], [plate()]),
                     (1, [code('MRSU4425761', score=0.95)], [plate()]),
                     (2, [code('MRSU4425761')], [plate()])])        # truck still in view
    assert len(events) == 1
    now, e = events[0]
    assert now == 1 and e['code']['code'] == 'MRSU4425761' and e['plate']['plate'] == '3231 UTA'
    assert e['code']['score'] == 0.95                                # best read is kept


def test_container_only_truck_is_logged_after_gap():
    t = GateTracker(min_hits=2, gap_s=5)
    events = run(t, [(0, [code('TRLU6689574')], []), (1, [code('TRLU6689574')], []),
                     (3, [], []), (7, [], [])])
    assert [(now, e['code']['code'], e['plate']) for now, e in events] == [(7, 'TRLU6689574', None)]


def test_plate_only_truck_after_max_wait():
    t = GateTracker(min_hits=2, max_wait_s=10)
    frames = [(s, [], [plate('5188', 'VDA')]) for s in range(12)]
    assert [now for now, _ in run(t, frames)] == [11]               # first confirmed at 1, +10 s


def test_cooldown_then_next_truck():
    t = GateTracker(min_hits=2, cooldown_s=60)
    first = run(t, [(0, [code('TRLU6689574')], [plate()]), (1, [code('TRLU6689574')], [plate()])])
    parked = run(t, [(s, [code('TRLU6689574')], [plate()]) for s in range(2, 30)])
    nxt = run(t, [(30, [code('TCNU6345807')], [plate('5489', 'EXA')]),
                  (31, [code('TCNU6345807')], [plate('5489', 'EXA')])])
    assert len(first) == 1 and parked == []
    assert [(e['code']['code'], e['plate']['plate']) for _, e in nxt] == [('TCNU6345807', '5489 EXA')]


def test_unverified_needs_double_votes():
    t = GateTracker(min_hits=2, gap_s=100, max_wait_s=100)
    assert run(t, [(i, [code('TCNU634580?', valid=False)], [plate(verified=False)]) for i in range(3)]) == []
    assert len(run(t, [(3, [code('TCNU634580?', valid=False)], [plate(verified=False)])])) == 1
    strict = GateTracker(min_hits=2, accept_unverified=False)
    assert run(strict, [(i, [code('TCNU634580?', valid=False)], []) for i in range(9)]) == []


def test_flush_emits_visit_in_progress():
    t = GateTracker(min_hits=2)
    run(t, [(0, [code('TRLU6689574')], []), (1, [code('TRLU6689574')], [])])
    assert [e['code']['code'] for e in t.flush(2)] == ['TRLU6689574']


def event(c=None, p=None):
    return {'code': c, 'plate': p, 'frame': FRAME, 'plate_frame': FRAME if p else None}


T0 = datetime(2026, 10, 15, 9, 0, 0)


def gate(tmp_path, **kw):
    return GateLog(Store(str(tmp_path / 'gate.db')), tmp_path / 'snaps', tmp_path / 'excel', **kw)


def test_rows_dedupe_merge_and_restart(tmp_path):
    log = gate(tmp_path)
    s = lambda sec: T0 + timedelta(seconds=sec)  # noqa: E731
    first = log.log(event(code('MRSU4425761')), direction='IN', when=s(0))                  # container only
    merged = log.log(event(code('MRSU4425761'), plate()), direction='IN', when=s(40))        # adds the plate
    dup = log.log(event(None, plate()), direction='IN', when=s(90))                          # same truck again
    out = log.log(event(code('MRSU4425761'), plate()), direction='OUT', when=s(120))         # leaving: new row
    later = log.log(event(code('MRSU4425761'), plate()), direction='IN', when=s(1800))       # back later: new row
    again = gate(tmp_path).log(event(code('MRSU4425761')), direction='IN', when=s(1860))     # after a restart
    assert [r['status'] for r in (first, merged, dup, out, later, again)] ==         ['added', 'merged', 'duplicate', 'added', 'added', 'duplicate']
    rows = log.store.day('2026-10-15')
    assert [(r['direction'], r['plate'], r['container']) for r in rows] ==         [('IN', '3231 UTA', 'MRSU 442576 1'), ('OUT', '3231 UTA', 'MRSU 442576 1'), ('IN', '3231 UTA', 'MRSU 442576 1')]
    assert rows[0]['plate_image'] and rows[0]['plate_verified'] == 1     # filled in by the merge
    # 3 rows x (snapshot + plate image); unused photos of merged/duplicate reads are deleted
    assert len(list((tmp_path / 'snaps').rglob('*.jpg'))) == 6


def test_unverified_container_upgraded_by_valid_read(tmp_path):
    log = gate(tmp_path)
    log.log(event(code('MRSU442576?', valid=False), plate()), direction='IN', when=T0)
    r = log.log(event(code('MRSU4425761')), direction='IN', when=T0 + timedelta(seconds=30))
    rows = log.store.day('2026-10-15')
    assert r['status'] == 'merged' and [(x['container'], x['check_digit_ok']) for x in rows] == [('MRSU 442576 1', 1)]


def test_excel_files_per_day_and_month(tmp_path):
    log = gate(tmp_path, language='ar')
    log.log(event(code('TRLU6689574'), plate()), direction='IN', when=datetime(2026, 10, 14, 23, 59))
    log.log(event(None, plate('5188', 'VDA')), direction='OUT', when=datetime(2026, 10, 15, 8, 0))
    assert log.flush() == set()
    day = load_workbook(tmp_path / 'excel' / 'gate_log_2026-10-15.xlsx').active
    month = list(load_workbook(tmp_path / 'excel' / 'gate_log_2026-10.xlsx').active.values)
    assert day.sheet_view.rightToLeft and list(day.values)[0][:3] == ('التاريخ', 'الوقت', 'الاتجاه')
    assert list(day.values)[1][2:4] == ('خروج', '5188 VDA') and len(list(day.values)) == 2
    assert [r[0] for r in month[1:]] == ['2026-10-14', '2026-10-15']


def test_excel_open_in_excel_is_retried_not_lost(tmp_path, monkeypatch):
    from pathlib import Path
    log = gate(tmp_path)
    log.log(event(code('TRLU6689574')), direction='IN', when=T0)
    real = Path.replace
    monkeypatch.setattr(Path, 'replace', lambda *a: (_ for _ in ()).throw(PermissionError))
    assert log.flush() == {'2026-10-15'}
    monkeypatch.setattr(Path, 'replace', real)
    assert log.flush() == set() and (tmp_path / 'excel' / 'gate_log_2026-10-15.xlsx').exists()
    assert len(log.store.day('2026-10-15')) == 1                          # data was never at risk


def test_correction_review_and_exports(tmp_path):
    log = gate(tmp_path)
    r = log.log(event(code('MRSU442576?', valid=False), plate(verified=False)), direction='IN', when=T0)
    row = log.store.get(r['id'])
    assert needs_review(row)
    log.store.correct(r['id'], container='MRSU 442576 1', check_digit_ok=1, plate_verified=1)
    row = log.store.get(r['id'])
    assert row['corrected'] == 1 and row['container'] == 'MRSU 442576 1' and not needs_review(row)
    xlsx = load_workbook(io.BytesIO(to_excel([row], 'en'))).active
    assert list(xlsx.values)[1][6:8] == ('MRSU 442576 1', 'YES')
    z = zipfile.ZipFile(io.BytesIO(to_zip([row], 'en', 'gate_log_2026-10-15')))
    names = sorted(z.namelist())
    assert names[0] == 'gate_log_2026-10-15.xlsx' and len(names) == 3 and all(n.startswith('photos/') for n in names[1:])
    link = load_workbook(io.BytesIO(z.read(names[0]))).active.cell(row=2, column=13).hyperlink.target
    assert link.startswith('photos/') and link in names                   # works after unzipping anywhere


def test_misreads_of_one_truck_combine_into_the_right_reading():
    # real reads of the MAERSK truck in the site footage (clip b, 57-59 s)
    t = GateTracker(min_hits=2, max_wait_s=100)
    reads = [code('IRSU1861395', valid=False), code('MRSU786395?', valid=False),
             code('MRSU7868395', valid=False), code('MRSU7860395', valid=False)]
    events = []
    for i, c in enumerate(reads):
        events += t.update([c], FRAME, i, [plate('231' if i == 0 else '3231', 'UTA', verified=False)])
    assert len(events) == 1
    e = events[0]
    assert (e['code']['code'], e['code']['valid']) == ('MRSU7869395', True)     # MRSU 786939 5
    assert e['plate']['plate'] == '3231 UTA'                                   # not the cut-off '231 UTA'


def test_a_plate_misread_that_flips_letters_does_not_make_a_second_truck():
    """One camera reads UTA, then UTR (letters differ, so a second plate group): still one truck."""
    t = GateTracker(min_hits=2)
    events = run(t, [(0, [], [plate()]), (0.5, [], [plate()]), (1, [], [plate(letters='UTR')]),
                     (1.5, [], [plate(letters='UTR')]), (2, [code('MRSU4425761')], []),
                     (2.5, [code('MRSU4425761')], [])]) + [(9, e) for e in t.flush(9)]
    assert [(e['code'] and e['code']['code'], e['plate']['plate']) for _, e in events] == [('MRSU4425761', '3231 UTA')]
