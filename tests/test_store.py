"""The gate records database: upgrading older databases, lane columns, Excel layout."""
import io
import sqlite3

from openpyxl import load_workbook

from src.i18n import t
from src.store import COLUMNS, LANE_COLUMNS, Store, needs_review, to_excel

# The trucks table as the previous version created it (before gate lanes)
OLD_SCHEMA = '''CREATE TABLE trucks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL, time TEXT NOT NULL, direction TEXT,
    plate TEXT, plate_ar TEXT, plate_verified INTEGER,
    container TEXT, check_digit_ok INTEGER, size_type TEXT,
    camera TEXT, location TEXT, confidence REAL,
    snapshot TEXT, plate_image TEXT, corrected INTEGER DEFAULT 0);
CREATE INDEX trucks_date ON trucks(date, time);'''

ROW = {'date': '2026-10-19', 'time': '09:00:00', 'direction': 'IN', 'plate': '3231 UTA', 'plate_verified': 1,
       'container': 'MRSU 442576 1', 'check_digit_ok': 1}


def test_older_database_is_upgraded_and_keeps_its_rows(tmp_path):
    path = tmp_path / 'gate.db'
    with sqlite3.connect(path) as db:
        db.executescript(OLD_SCHEMA)
        db.execute("INSERT INTO trucks (date, time, direction, plate, snapshot) VALUES "
                   "('2026-10-18', '08:00:00', 'OUT', '5188 VDA', 'C:/snaps/a.jpg')")
    store = Store(str(path))
    [old] = store.day('2026-10-18')
    assert (old['plate'], old['snapshot'], old['lane'], old['cams_agree']) == ('5188 VDA', 'C:/snaps/a.jpg', None, None)
    _, rid = store.add({**ROW, 'lane': 'Entry', 'cams_agree': 2, 'cams_expected': 2, 'plate_source': 'anpr',
                        'snapshot2': 'C:/snaps/b.jpg'})
    new = store.get(rid)
    assert (new['lane'], new['cams_agree'], new['plate_source'], new['snapshot2']) == ('Entry', 2, 'anpr', 'C:/snaps/b.jpg')
    Store(str(path))  # opening an upgraded database again changes nothing


def test_single_camera_sites_keep_todays_excel_layout(tmp_path):
    ws = load_workbook(io.BytesIO(to_excel([ROW], 'en'))).active
    assert [c.value for c in ws[1]] == [t(c, 'en') for c in COLUMNS]


def test_two_camera_lanes_get_verified_and_second_photo_columns(tmp_path):
    snap2 = tmp_path / 'b.jpg'
    rows = [{**ROW, 'cams_agree': 2, 'cams_expected': 2, 'snapshot2': str(snap2)},
            {**ROW, 'time': '09:05:00', 'cams_agree': 1, 'cams_expected': 2},
            {**ROW, 'time': '09:10:00', 'container': None}]
    ws = load_workbook(io.BytesIO(to_excel(rows, 'en'))).active
    head = [c.value for c in ws[1]]
    assert head == [t(c, 'en') for c in LANE_COLUMNS]
    col = head.index('Container cameras')
    assert [ws.cell(row=r, column=col + 1).value for r in (2, 3, 4)] == ['Verified by 2 cameras', '1 camera only', None]
    assert ws.cell(row=2, column=len(head)).hyperlink.target == str(snap2)
    ar = load_workbook(io.BytesIO(to_excel(rows, 'ar'))).active
    assert ar.cell(row=3, column=col + 1).value == 'كاميرا واحدة فقط'


def test_needs_review_when_only_one_of_two_cameras_agreed():
    assert not needs_review(ROW)                                              # single camera, all confirmed
    assert not needs_review({**ROW, 'cams_agree': 2, 'cams_expected': 2})
    assert needs_review({**ROW, 'cams_agree': 1, 'cams_expected': 2})
    assert not needs_review({**ROW, 'cams_agree': 1, 'cams_expected': 2, 'corrected': 1})  # a person checked it
    assert not needs_review({**ROW, 'cams_agree': 1, 'cams_expected': 1})


def test_merge_carries_lane_fields_and_corrections_leave_them_alone(tmp_path):
    store = Store(str(tmp_path / 'gate.db'))
    _, rid = store.add({**ROW, 'container': None, 'check_digit_ok': None, 'plate_source': 'anpr'})
    status, same = store.add({**ROW, 'time': '09:00:20', 'cams_agree': 2, 'cams_expected': 2})
    row = store.get(rid)
    assert (status, same) == ('merged', rid)
    assert (row['container'], row['cams_agree'], row['cams_expected'], row['plate_source']) == ('MRSU 442576 1', 2, 2, 'anpr')
    store.correct(rid, plate='3231 UTA', cams_agree=9, snapshot2='C:/elsewhere.jpg', lane='X')
    row = store.get(rid)
    assert (row['cams_agree'], row['snapshot2'], row['lane'], row['corrected']) == (2, None, None, 1)
