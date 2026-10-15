"""Turn per-frame reads into one gate event per truck (plate + container) and log it to
Excel with snapshots."""
import io
import threading
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import cv2
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.i18n import t
from src.iso6346 import format_code
from src.recognizer import draw

COLUMNS = ['date', 'time', 'direction', 'plate', 'plate_ar', 'plate_check', 'container',
           'check_digit', 'size_type', 'camera', 'location', 'confidence', 'snapshot', 'plate_image']
LINKS = ('snapshot', 'plate_image')


class _Votes:
    """Reads of one kind (container codes or plates) during a visit."""

    def __init__(self):
        self.hits = defaultdict(list)  # key -> [timestamps]
        self.best = {}                 # key -> (score, item, frame)

    def add(self, key, item, frame, now, window_s):
        self.hits[key] = [t for t in self.hits[key] if now - t <= window_s] + [now]
        if key not in self.best or item['score'] > self.best[key][0]:
            self.best[key] = (item['score'], item, frame.copy())

    def confirmed(self, need):
        """Most-read key that has enough reads (need(item) -> required count), or None."""
        ok = [k for k in self.hits if len(self.hits[k]) >= need(self.best[k][1])]
        return max(ok, key=lambda k: (len(self.hits[k]), self.best[k][0])) if ok else None


class GateTracker:
    """A truck passing a camera is seen in many frames. Plate and container code are each
    confirmed once read `min_hits` times within `window_s` (2x for an unconfirmed check digit
    or a plate whose Arabic row didn't match). One event is emitted per truck visit:
    - as soon as both are confirmed, or
    - with whatever was confirmed after `max_wait_s`, or when nothing new has been read for
      `gap_s` (truck left, or one without container / with an unreadable plate).
    Logged codes and plates are then ignored for `cooldown_s` (truck waiting at the barrier).
    Call update() for every processed frame, also empty ones, so visits can end."""

    def __init__(self, min_hits=2, window_s=15, cooldown_s=300, accept_unverified=True,
                 gap_s=5, max_wait_s=10):
        self.min_hits, self.window_s, self.cooldown_s = min_hits, window_s, cooldown_s
        self.accept_unverified, self.gap_s, self.max_wait_s = accept_unverified, gap_s, max_wait_s
        self.logged = {}   # 'serial' or 'plate' key -> last event time
        self.visit = None

    def update(self, codes, frame, now=None, plates=()):
        """Feed one frame's reads; returns a list of events. An event is a dict with
        code (or None), plate (or None), frame, plate_frame."""
        now = time.time() if now is None else now
        codes = [c for c in codes if (c['valid'] or self.accept_unverified)
                 and not self._cooling(c['code'][4:10], now)]
        plates = [p for p in plates if not self._cooling(p['digits'] + p['letters'], now)]
        events = []
        v = self.visit
        if v and now - v['last'] > self.gap_s:  # visit over
            events += self._emit(now)
            v = None
        if codes or plates:
            if v is None:
                v = self.visit = {'codes': _Votes(), 'plates': _Votes(), 'last': now, 'first_ok': None}
            v['last'] = now
            for c in codes:
                v['codes'].add(c['code'], c, frame, now, self.window_s)
            for p in plates:
                v['plates'].add(p['digits'] + p['letters'], p, frame, now, self.window_s)
        if v:
            code, plate = self._confirmed()
            if code or plate:
                v['first_ok'] = v['first_ok'] or now
                if (code and plate) or now - v['first_ok'] >= self.max_wait_s:
                    events += self._emit(now)
        return events

    def flush(self, now=None):
        """Emit the visit in progress (end of a video file)."""
        return self._emit(time.time() if now is None else now)

    def _cooling(self, key, now):
        return now - self.logged.get(key, -1e18) < self.cooldown_s

    def _confirmed(self):
        v, n = self.visit, self.min_hits
        code = v['codes'].confirmed(lambda c: n if c['valid'] else 2 * n)
        plate = v['plates'].confirmed(lambda p: n if p['verified'] else 2 * n)
        return code, plate

    def _emit(self, now):
        if not self.visit:
            return []
        code, plate = self._confirmed()
        v, self.visit = self.visit, None
        if not (code or plate):
            return []
        c = v['codes'].best[code] if code else None
        p = v['plates'].best[plate] if plate else None
        if code:
            self.logged[code[4:10]] = now
        if plate:
            self.logged[plate] = now
        return [{'code': c and c[1], 'plate': p and p[1],
                 'frame': (c or p)[2], 'plate_frame': p and p[2]}]


class ExcelLog:
    """Append-only Excel log, one row per truck. `path` may contain strftime codes, e.g.
    logs/gate_%Y-%m.xlsx for one workbook per month. language 'en' or 'ar' sets headers and
    values (Arabic sheets are right-to-left). Thread-safe. If the workbook is open in Excel
    (locked on Windows), rows are queued and written on the next append/flush. A workbook
    with different columns (older version, other language) is left alone and a new one
    with a _2, _3... suffix is used.

    Duplicates are checked against the workbook itself, so they are caught across restarts,
    cameras and processes: a row for the same direction and the same container (serial) or
    plate within `dedupe_s` seconds is not added again. If the new reading has something the
    earlier row lacks (a plate, a container, a confirmed check digit), that row is filled in
    instead. row['status'] says what happened: 'added', 'merged', 'duplicate' or 'queued'."""

    def __init__(self, path='logs/gate_log_%Y-%m.xlsx', snapshot_dir='logs/snapshots', language='en',
                 dedupe_s=300):
        self.path, self.snapshot_dir, self.lang, self.dedupe_s = path, Path(snapshot_dir), language, dedupe_s
        self.pending = []
        self.lock = threading.Lock()

    def log(self, event, camera='camera', location='', direction='', when=None):
        """Save the snapshots and append one row. Returns the row as {column key: value}."""
        when = when or datetime.now()
        code, plate = event['code'], event['plate']
        day_dir = self.snapshot_dir / when.strftime('%Y-%m-%d')
        day_dir.mkdir(parents=True, exist_ok=True)
        name = '_'.join(x for x in (code and code['code'].rstrip('?'),
                                    plate and plate['digits'] + plate['letters']) if x)
        snap = day_dir / f"{when:%H%M%S}_{_safe(camera)}_{name}.jpg"
        cv2.imwrite(str(snap), draw(event['frame'].copy(), [code] if code else [], [plate] if plate else []))
        plate_img = ''
        if plate:
            x, y, w, h = plate['box']
            m = max(4, h // 4)
            crop = event['plate_frame'][max(0, y - m):y + h + m, max(0, x - m):x + w + m]
            if crop.size:
                plate_img = day_dir / f"{when:%H%M%S}_{_safe(camera)}_{name}_plate.jpg"
                f = max(1.0, 120 / crop.shape[0])
                cv2.imwrite(str(plate_img), cv2.resize(crop, None, fx=f, fy=f, interpolation=cv2.INTER_CUBIC))
                plate_img = str(plate_img.resolve())
        L = self.lang
        row = {
            'date': when.strftime('%Y-%m-%d'), 'time': when.strftime('%H:%M:%S'),
            'direction': t(direction, L) if direction else '',
            'plate': plate['plate'] if plate else '',
            'plate_ar': plate['arabic'] if plate else '',
            'plate_check': t('verified' if plate['verified'] else 'latin_only', L) if plate else '',
            'container': format_code(code['code']) if code else '',
            'check_digit': t('yes' if code['valid'] else 'no', L) if code else '',
            'size_type': (code['size_type'] or '') if code else '',
            'camera': camera, 'location': location,
            'confidence': round(min(x['score'] for x in (code, plate) if x), 3),
            'snapshot': str(snap.resolve()), 'plate_image': plate_img, 'status': 'queued',
        }
        with self.lock:
            self.pending.append((when, row))
            self._flush()
        return row

    def flush(self):
        with self.lock:
            self._flush()

    def _workbook(self, path):
        """Open (or create) the workbook at path; on a column mismatch try path_2, path_3..."""
        headers = [t(c, self.lang) for c in COLUMNS]
        base, n = path, 1
        while path.exists():
            wb = load_workbook(path)
            if [c.value for c in wb.active[1]] == headers:
                return path, wb
            n += 1
            path = base.with_name(f"{base.stem}_{n}{base.suffix}")
        path.parent.mkdir(parents=True, exist_ok=True)
        return path, _new_workbook(self.lang)

    def _flush(self):
        while self.pending:
            when, row = self.pending[0]
            try:
                path, wb = self._workbook(Path(when.strftime(self.path)))
                ws = wb.active
                r = self._match(ws, row, when)
                if r is None:
                    ws.append([row[c] for c in COLUMNS])
                    for key in LINKS:
                        self._set(ws, ws.max_row, key, row[key])
                    status = 'added'
                else:
                    status = 'merged' if self._merge(ws, r, row) else 'duplicate'
                if status != 'duplicate':
                    wb.save(path)
            except PermissionError:
                print(f"[gate] {when.strftime(self.path)} is open in another program, {len(self.pending)} "
                      f"row(s) queued; close it and they will be written on the next event.")
                return
            row['status'] = status
            if status != 'added':  # the earlier row keeps its snapshot; drop the unused new files
                for key in LINKS:
                    if row[key] and not (status == 'merged' and key == 'plate_image' and row.get('_plate_used')):
                        Path(row[key]).unlink(missing_ok=True)
            self.pending.pop(0)

    def _set(self, ws, r, key, value):
        _set_cell(ws, r, key, value, self.lang)

    def _match(self, ws, row, when):
        """Row number of a recent row for the same truck, or None."""
        if not self.dedupe_s:
            return None
        i = {c: n for n, c in enumerate(COLUMNS)}
        serial = _serial(row['container'])
        for r in range(ws.max_row, max(1, ws.max_row - 300), -1):  # newest first
            v = [c.value for c in ws[r]]
            try:
                logged = datetime.strptime(f"{v[i['date']]} {v[i['time']]}", '%Y-%m-%d %H:%M:%S')
            except (TypeError, ValueError):
                continue
            if (when - logged).total_seconds() > self.dedupe_s:
                break  # rows are in time order: everything above is older
            if (v[i['direction']] or '') != row['direction']:
                continue
            if (serial and serial == _serial(v[i['container']])) or (row['plate'] and row['plate'] == v[i['plate']]):
                return r
        return None

    def _merge(self, ws, r, row):
        """Fill in what the earlier row of the same truck is missing. True if it changed."""
        get = lambda k: ws.cell(row=r, column=COLUMNS.index(k) + 1).value  # noqa: E731
        changed = False
        if row['plate'] and not get('plate'):
            for k in ('plate', 'plate_ar', 'plate_check', 'plate_image'):
                self._set(ws, r, k, row[k])
            row['_plate_used'] = changed = True
        if row['container'] and (not get('container') or (
                get('check_digit') == t('no', self.lang) and row['check_digit'] == t('yes', self.lang))):
            for k in ('container', 'check_digit', 'size_type'):
                self._set(ws, r, k, row[k])
            changed = True
        return changed


def _new_workbook(lang):
    wb = Workbook()
    ws = wb.active
    ws.title = t('sheet', lang)
    ws.sheet_view.rightToLeft = lang == 'ar'
    ws.append([t(c, lang) for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    return wb


def _set_cell(ws, r, key, value, lang):
    cell = ws.cell(row=r, column=COLUMNS.index(key) + 1)
    if key in LINKS:
        if value:
            cell.hyperlink, cell.value, cell.style = value, t('open_image', lang), 'Hyperlink'
    else:
        cell.value = value


def read_day(log_dir, day):
    """All rows logged on `day` ('YYYY-MM-DD') in the gate workbooks in log_dir, oldest
    first, as {column key: value} (link columns hold the image file path). Workbooks in
    either language are read; ones with other columns (older versions) are skipped."""
    headers = [[t(c, lang) for c in COLUMNS] for lang in ('en', 'ar')]
    rows = []
    for path in sorted(Path(log_dir).glob('*.xlsx')):
        try:
            ws = load_workbook(path).active
        except Exception:  # being written right now, or not a gate log
            continue
        if [c.value for c in ws[1]] not in headers:
            continue
        for cells in ws.iter_rows(min_row=2, max_col=len(COLUMNS)):
            if str(cells[0].value) == day:
                rows.append({k: (c.hyperlink.target if c.hyperlink else '') if k in LINKS else c.value
                             for k, c in zip(COLUMNS, cells)})
    return sorted(rows, key=lambda r: str(r['time']))


def export_day(rows, lang='en'):
    """One day's rows (from read_day) as .xlsx file bytes, same layout as the gate log."""
    wb = _new_workbook(lang)
    ws = wb.active
    for row in rows:
        ws.append([None if k in LINKS else row[k] for k in COLUMNS])
        for key in LINKS:
            _set_cell(ws, ws.max_row, key, row[key], lang)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _serial(container):
    """'MRSU 442576 1' -> 'MRSU442576' (owner + serial; the check digit may differ between reads)."""
    c = ''.join((container or '').split())
    return c[:10] if len(c) >= 10 else ''


def _safe(name):
    return ''.join(c if c.isalnum() or c in '-_' else '_' for c in name)
