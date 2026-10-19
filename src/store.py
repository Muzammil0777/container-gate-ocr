"""Gate records: an SQLite database is the master record (safe with several readers and
writers, survives power cuts); Excel files (one per day, one per month) are generated from it.
"""
import io
import sqlite3
import threading
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.i18n import t
from src.iso6346 import check_digit, format_code, normalize
from src.saudi_plate import format_plate, parse_latin, to_arabic

# Gate lanes (several cameras per truck): added to existing databases on start
LANE_FIELDS = {'lane': 'TEXT',
               'cams_agree': 'INTEGER',     # container cameras whose reading agreed with the row's
               'cams_expected': 'INTEGER',  # container cameras in the lane (2 = can be verified)
               'plate_source': 'TEXT',      # 'anpr' (the plate camera's own reading) or 'ocr'
               'snapshot2': 'TEXT'}         # the second container camera's photo
FIELDS = ['date', 'time', 'direction', 'plate', 'plate_ar', 'plate_verified', 'container',
          'check_digit_ok', 'size_type', 'camera', 'location', 'confidence', 'snapshot',
          'plate_image', 'corrected', *LANE_FIELDS]
# Excel columns (same order as before) and how each value is shown
COLUMNS = ['date', 'time', 'direction', 'plate', 'plate_ar', 'plate_check', 'container',
           'check_digit', 'size_type', 'camera', 'location', 'confidence', 'snapshot', 'plate_image']
# a site with two container cameras in a lane gets two more: verified, and the second photo
LANE_COLUMNS = COLUMNS[:8] + ['cams_verified'] + COLUMNS[8:] + ['snapshot2']
LINKS = ('snapshot', 'plate_image', 'snapshot2')
WIDTHS = {'date': 11, 'time': 9, 'direction': 9, 'plate': 11, 'plate_ar': 14, 'plate_check': 20, 'container': 16,
          'check_digit': 10, 'cams_verified': 20, 'size_type': 9, 'camera': 16, 'location': 18, 'confidence': 10}

SCHEMA = '''CREATE TABLE IF NOT EXISTS trucks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    date TEXT NOT NULL, time TEXT NOT NULL,      -- YYYY-MM-DD, HH:MM:SS (local time)
    direction TEXT,                               -- IN / OUT
    plate TEXT, plate_ar TEXT, plate_verified INTEGER,
    container TEXT, check_digit_ok INTEGER, size_type TEXT,
    camera TEXT, location TEXT, confidence REAL,
    snapshot TEXT, plate_image TEXT,
    corrected INTEGER DEFAULT 0                   -- 1 once a person fixed the reading
);
CREATE INDEX IF NOT EXISTS trucks_date ON trucks(date, time);
CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT);  -- app state (license, clock guard)'''


def _serial(container):
    """'MRSU 442576 1' -> 'MRSU442576' (owner + serial; the check digit may differ between reads)."""
    c = ''.join((container or '').split())
    return c[:10] if len(c) >= 10 else ''


class Store:
    """All gate records. Thread-safe; one instance per process is enough."""

    def __init__(self, path='logs/gate.db'):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.db = sqlite3.connect(path, check_same_thread=False, timeout=30)
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA journal_mode=WAL')  # readers never block the writer
        self.db.executescript(SCHEMA)
        have = {r['name'] for r in self.db.execute('PRAGMA table_info(trucks)')}
        for col, kind in LANE_FIELDS.items():  # an older database: add the new columns, keep every row
            if col not in have:
                self.db.execute(f'ALTER TABLE trucks ADD COLUMN {col} {kind}')
        self.lock = threading.Lock()

    def add(self, row, dedupe_s=300):
        """Insert a truck, unless the same truck (same direction, same container serial or
        plate within dedupe_s seconds) is already there: then fill in what that row is
        missing. Returns (status, id) with status 'added', 'merged' or 'duplicate'."""
        with self.lock, self.db:
            prev = self._recent_match(row, dedupe_s) if dedupe_s else None
            if prev is None:
                cur = self.db.execute(f"INSERT INTO trucks ({', '.join(FIELDS)}) VALUES ({', '.join('?' * len(FIELDS))})",
                                      [row.get(f) for f in FIELDS])
                return 'added', cur.lastrowid
            updates = {}
            if row.get('plate') and not prev['plate']:
                updates.update({k: row.get(k) for k in ('plate', 'plate_ar', 'plate_verified', 'plate_image', 'plate_source')})
            if row.get('container') and (not prev['container'] or (not prev['check_digit_ok'] and row.get('check_digit_ok'))):
                updates.update({k: row.get(k) for k in ('container', 'check_digit_ok', 'size_type', 'cams_agree', 'cams_expected')})
            if not updates:
                return 'duplicate', prev['id']
            self._update(prev['id'], updates)
            return 'merged', prev['id']

    def _recent_match(self, row, dedupe_s):
        when = datetime.strptime(f"{row['date']} {row['time']}", '%Y-%m-%d %H:%M:%S')
        since = when - timedelta(seconds=dedupe_s)
        serial, plate = _serial(row.get('container')), row.get('plate') or ''
        cands = self.db.execute(
            "SELECT * FROM trucks WHERE direction IS ? AND (date || ' ' || time) BETWEEN ? AND ? ORDER BY id DESC",
            (row.get('direction'), since.strftime('%Y-%m-%d %H:%M:%S'), when.strftime('%Y-%m-%d %H:%M:%S')))
        for r in cands:
            if (serial and serial == _serial(r['container'])) or (plate and plate == r['plate']):
                return r
        return None

    def _update(self, rid, updates):
        self.db.execute(f"UPDATE trucks SET {', '.join(f'{k} = ?' for k in updates)} WHERE id = ?",
                        [*updates.values(), rid])

    def correct(self, rid, **fields):
        """A person fixed a reading (plate, plate_ar, container, ...): save it and mark the row."""
        allowed = {k: v for k, v in fields.items() if k in FIELDS and k not in ('date', 'time', *LINKS, *LANE_FIELDS)}
        with self.lock, self.db:
            self._update(rid, {**allowed, 'corrected': 1})

    def meta_get(self, k):
        r = self.db.execute('SELECT v FROM meta WHERE k = ?', (k,)).fetchone()
        return r[0] if r else None

    def meta_set(self, k, v):
        with self.lock, self.db:
            self.db.execute('INSERT INTO meta (k, v) VALUES (?, ?) ON CONFLICT(k) DO UPDATE SET v = excluded.v', (k, v))

    def get(self, rid):
        r = self.db.execute('SELECT * FROM trucks WHERE id = ?', (rid,)).fetchone()
        return dict(r) if r else None

    def day(self, day):
        """Rows of one day ('YYYY-MM-DD'), oldest first, as dicts (with 'id')."""
        return [dict(r) for r in self.db.execute('SELECT * FROM trucks WHERE date = ? ORDER BY time, id', (day,))]

    def month(self, month):
        """Rows of one month ('YYYY-MM')."""
        return [dict(r) for r in self.db.execute("SELECT * FROM trucks WHERE date LIKE ? ORDER BY date, time, id",
                                                 (month + '-%',))]

    def months(self):
        return [r[0] for r in self.db.execute("SELECT DISTINCT substr(date, 1, 7) FROM trucks ORDER BY 1 DESC")]


def parse_fix(plate, container, direction):
    """A person's typed correction -> (fields for Store.correct, error text keys). Empty text
    clears that reading. A container whose check digit doesn't match is still accepted
    (fields['check_digit_ok'] == 0), the caller should warn."""
    fix, errors = {'direction': direction}, []
    if plate.strip():
        p = parse_latin(plate)
        if p:
            fix.update(plate=format_plate(*p), plate_ar=to_arabic(*p), plate_verified=1)
        else:
            errors.append('bad_plate')
    else:
        fix.update(plate=None, plate_ar=None, plate_verified=None)
    if container.strip():
        c = normalize(container)
        if len(c) == 11 and c[:4].isalpha() and c[4:].isdigit():
            fix.update(container=format_code(c), check_digit_ok=int(check_digit(c[:10]) == int(c[10])))
        else:
            errors.append('bad_container')
    else:
        fix.update(container=None, check_digit_ok=None, size_type=None)
    return fix, errors


def needs_review(row):
    """Readings worth a second look: plate not confirmed by its Arabic row, container check
    digit not confirmed, either one missing, or (in a lane with two container cameras) the
    container read by only one of them. Rows a person already fixed are fine."""
    if row.get('corrected'):
        return False
    return (not row.get('plate') or not row.get('plate_verified')
            or not row.get('container') or not row.get('check_digit_ok') or not cams_verified(row))


def cams_verified(row):
    """False only when fewer of the lane's container cameras agreed on the container than it has."""
    expected = row.get('cams_expected') or 0
    return expected < 2 or (row.get('cams_agree') or 0) >= expected


def display(row, lang='en'):
    """A stored row as Excel/table values in the chosen language, keyed by COLUMNS."""
    yes_no = lambda v: '' if v is None else t('yes' if v else 'no', lang)  # noqa: E731
    return {
        'date': row['date'], 'time': row['time'],
        'direction': t(row['direction'], lang) if row.get('direction') else '',
        'plate': row.get('plate') or '', 'plate_ar': row.get('plate_ar') or '',
        'plate_check': '' if not row.get('plate') else t('verified' if row.get('plate_verified') else 'latin_only', lang),
        'container': row.get('container') or '',
        'check_digit': '' if not row.get('container') else yes_no(row.get('check_digit_ok')),
        'size_type': row.get('size_type') or '', 'camera': row.get('camera') or '',
        'location': row.get('location') or '', 'confidence': row.get('confidence'),
        'cams_verified': '' if not row.get('container') or (row.get('cams_expected') or 0) < 2
        else t('cams_agreed', lang, n=row.get('cams_agree')) if cams_verified(row) else t('one_cam_only', lang),
        'snapshot': row.get('snapshot') or '', 'plate_image': row.get('plate_image') or '',
        'snapshot2': row.get('snapshot2') or '',
    }


def to_excel(rows, lang='en', links=None):
    """Rows as .xlsx bytes in the gate-log layout. links maps each LINKS column to a function
    turning the stored image path into the hyperlink target (default: the path itself)."""
    columns = LANE_COLUMNS if any((r.get('cams_expected') or 0) >= 2 for r in rows) else COLUMNS
    wb = Workbook()
    ws = wb.active
    ws.title = t('sheet', lang)
    ws.sheet_view.rightToLeft = lang == 'ar'
    ws.append([t(c, lang) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for row in rows:
        d = display(row, lang)
        ws.append([None if c in LINKS else d[c] for c in columns])
        for key in LINKS:
            if key in columns and d[key]:
                cell = ws.cell(row=ws.max_row, column=columns.index(key) + 1)
                target = (links or {}).get(key, lambda p: p)(d[key])
                cell.hyperlink, cell.value, cell.style = target, t('open_image', lang), 'Hyperlink'
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 12)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_zip(rows, lang='en', name='gate_log'):
    """Excel plus every truck's photos as one .zip; the Excel links point into its photos/
    folder, so they work on any computer once the zip is extracted."""
    rel = lambda p: f"photos/{Path(p).name}"  # noqa: E731
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr(f'{name}.xlsx', to_excel(rows, lang, {k: rel for k in LINKS}))
        for row in rows:
            for key in LINKS:
                if row.get(key) and Path(row[key]).is_file():
                    z.write(row[key], rel(row[key]))
    return buf.getvalue()
